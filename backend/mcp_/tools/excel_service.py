from __future__ import annotations

from typing import Any

from mcp.server.fastmcp import FastMCP
from openpyxl import load_workbook

from utils.decorators import _tool
from utils.logger import logger


def register_excel_tools(mcp: FastMCP) -> None:
    tool = _tool(mcp)

    @tool
    def update_excel_cell(file_path: str, row: int, column: str, value: str) -> dict[str, Any]:
        """Update a single cell in an Excel file with a string value.

        Args:
            file_path: Path to the .xlsx file to modify.
            row: 1-based row number of the target cell.
            column: Column letter of the target cell (A, B, C, ...).
            value: String value to write into the cell.

        Returns:
            dict with "success", "cell" (e.g. "B3"), and "value".
        """
        wb = load_workbook(file_path)
        ws = wb.active
        ws[f"{column.upper()}{row}"].value = value
        wb.save(file_path)
        logger.info(f"Updated cell {column.upper()}{row} with value {value}")
        wb.close()
        return {"success": True, "cell": f"{column.upper()}{row}", "value": value}

    @tool
    def read_excel_cell(file_path: str, row: int, column: str) -> dict[str, Any]:
        """Read the value of a single cell from an Excel file.

        Args:
            file_path: Path to the .xlsx file to read.
            row: 1-based row number of the target cell.
            column: Column letter of the target cell (A, B, C, ...).

        Returns:
            dict with "cell" (e.g. "B3") and "value" (string or null when empty).
        """
        wb = load_workbook(file_path, read_only=True)
        ws = wb.active
        value = ws[f"{column.upper()}{row}"].value
        wb.close()
        return {"cell": f"{column.upper()}{row}", "value": str(value) if value is not None else None}
