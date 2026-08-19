from __future__ import annotations

from datetime import datetime, timezone
from typing import Any

from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from utils.logger import logger

REQUIRED_COLUMNS = ["post_type", "scheduled_time"]

ALLOWED_POST_TYPES: dict[str, list[str]] = {
    "instagram_image":    ["media_url", "message"],
    "instagram_reel":     ["media_url", "message"],
    "instagram_story":    ["media_url"],
    "instagram_carousel": ["media_url", "message"],
    "whatsapp_text":      ["recipient_id", "message"],
    "whatsapp_image":     ["recipient_id", "media_url", "message"],
    "whatsapp_video":     ["recipient_id", "media_url", "message"],
    "facebook_post":      ["message", "media_url"],
}


def _load_workbook(file_path: str) -> tuple[Any, list[str], list[Any]]:
    try:
        wb = load_workbook(file_path, read_only=True)
    except (InvalidFileException, BadZipFile) as e:
        raise ValueError("File is not a valid .xlsx Excel file") from e
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        wb.close()
        raise ValueError("Excel file is empty")

    headers = [str(h).lower().strip() if h else "" for h in rows[0]]
    for col in REQUIRED_COLUMNS:
        if col not in headers:
            wb.close()
            raise ValueError(f"Excel must contain column: {col}. Found: {headers}")

    return wb, headers, rows[1:]


def _parse_row_to_entry(row: tuple, headers: list[str]) -> dict[str, Any]:
    entry = {}
    for j, header in enumerate(headers):
        entry[header] = row[j] if j < len(row) else None
    return entry


def _validate_post_type(entry: dict) -> str | None:
    post_type = str(entry.get("post_type", "")).strip().lower()
    if not post_type:
        return None
    if post_type not in ALLOWED_POST_TYPES:
        return None
    return post_type


def _normalize_scheduled_time(entry: dict) -> datetime | None:
    val = entry.get("scheduled_time")
    if isinstance(val, str):
        try:
            val = datetime.fromisoformat(val)
        except ValueError:
            return None

    if not isinstance(val, datetime):
        return None

    if val.tzinfo is not None:
        val = val.astimezone(timezone.utc).replace(tzinfo=None)

    if val < datetime.now(timezone.utc).replace(tzinfo=None):
        return None

    return val


def _build_post(entry: dict, post_type: str, scheduled_time: datetime) -> dict[str, Any]:
    return {
        "post_type": post_type,
        "platform": str(entry.get("platform")).strip(),
        "media_url": str(entry.get("media_url")).strip(),
        "message": str(entry.get("message")).strip(),
        "topic": str(entry.get("topic")).strip(),
        "recipient_id": str(entry.get("recipient_id")).strip(),
        #"location_id": str(entry.get("location_id")).strip(),
        "scheduled_time": scheduled_time,
    }


def parse_excel(file_path: str) -> list[dict[str, Any]]:
    logger.info("Parsing Excel file | path=%s", file_path)
    wb, headers, rows = _load_workbook(file_path)
    posts: list[dict[str, Any]] = []
    errors: list[str] = []

    for i, row in enumerate(rows, start=2):
        if all(cell is None for cell in row):
            continue

        entry = _parse_row_to_entry(row, headers)

        post_type = _validate_post_type(entry)
        if post_type is None:
            pt = str(entry.get("post_type", "")).strip()
            msg = f"Row {i}: {'missing' if not pt else 'invalid'} post_type '{pt}'"
            errors.append(msg)
            continue

        scheduled_time = _normalize_scheduled_time(entry)
        if scheduled_time is None:
            errors.append(f"Row {i}: invalid or past scheduled_time")
            continue

        posts.append(_build_post(entry, post_type, scheduled_time))

    wb.close()

    if errors:
        logger.error("Excel parsing errors:\n%s", "\n".join(errors))
        raise ValueError("Excel parsing errors:\n" + "\n".join(errors))

    logger.info("Excel parsed | rows=%d posts=%d", len(rows), len(posts))
    return posts
